import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_HEADER_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_FILL  = PatternFill("solid", fgColor="1E1E2E")
_CENTER       = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT         = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_THIN         = Side(style="thin", color="DDDDDD")
_BORDER       = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_TEMP_FILL = {
    "HOT":  PatternFill("solid", fgColor="FFE5E5"),
    "WARM": PatternFill("solid", fgColor="FFFBE5"),
    "COLD": PatternFill("solid", fgColor="E5F0FF"),
}
_TEMP_FONT = {
    "HOT":  Font(name="Calibri", color="CC0000", bold=True, size=10),
    "WARM": Font(name="Calibri", color="AA7700", bold=True, size=10),
    "COLD": Font(name="Calibri", color="0055AA", bold=True, size=10),
}


def _set_header(ws, col: int, label: str, width: float):
    cell = ws.cell(row=1, column=col, value=label)
    cell.font   = _HEADER_FONT
    cell.fill   = _HEADER_FILL
    cell.alignment = _CENTER
    cell.border = _BORDER
    ws.column_dimensions[get_column_letter(col)].width = width


def build(leads: list[dict]) -> bytes:
    """Return a formatted .xlsx as raw bytes."""
    wb = Workbook()

    # ── Sheet 1: All Leads ───────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "All Leads"
    ws1.sheet_view.showGridLines = False
    ws1.freeze_panes = "A2"

    cols1 = [
        ("Name", 22), ("Email", 28), ("Company", 22), ("Job Title", 20),
        ("Phone", 16), ("Temperature", 13), ("Next Step", 34),
        ("Notes", 38), ("Captured By", 16), ("Captured At", 20),
    ]
    for i, (label, width) in enumerate(cols1, 1):
        _set_header(ws1, i, label, width)
    ws1.row_dimensions[1].height = 30

    for row_i, lead in enumerate(leads, 2):
        temp = lead.get("temperature", "WARM")
        fill = _TEMP_FILL.get(temp)
        ts = lead.get("captured_at", "")
        if hasattr(ts, "strftime"):
            ts = ts.strftime("%Y-%m-%d %H:%M")

        values = [
            lead.get("name", ""),      lead.get("email", ""),
            lead.get("company", ""),   lead.get("job_title", ""),
            lead.get("phone", ""),     temp,
            lead.get("next_step", ""), lead.get("notes", ""),
            lead.get("captured_by", ""), ts,
        ]
        for col_i, val in enumerate(values, 1):
            cell = ws1.cell(row=row_i, column=col_i, value=val)
            cell.border = _BORDER
            cell.alignment = _CENTER if col_i == 6 else _LEFT
            if fill:
                cell.fill = fill
            if col_i == 6:
                cell.font = _TEMP_FONT.get(temp, Font(name="Calibri", size=10))
        ws1.row_dimensions[row_i].height = 20

    # Summary block
    if leads:
        sr = len(leads) + 3
        counts = {"HOT": 0, "WARM": 0, "COLD": 0}
        for r in leads:
            t = r.get("temperature", "WARM")
            counts[t] = counts.get(t, 0) + 1
        ws1.cell(row=sr,     column=1, value="Summary").font = Font(bold=True)
        ws1.cell(row=sr,     column=2, value=f"Total: {len(leads)}")
        for offset, (t, label) in enumerate([("HOT","🔥 HOT"),("WARM","🌤 WARM"),("COLD","❄️ COLD")], 1):
            ws1.cell(row=sr + offset, column=1, value=label).font = _TEMP_FONT[t]
            ws1.cell(row=sr + offset, column=2, value=counts[t])

    # ── Sheet 2: Pipedrive Import ────────────────────────────────────────────
    ws2 = wb.create_sheet("Pipedrive Import")
    ws2.sheet_view.showGridLines = False
    ws2.freeze_panes = "A2"

    cols2 = [
        ("Name", 22), ("Email", 28), ("Phone", 16),
        ("Organization name", 22), ("Job title", 20), ("Notes", 55),
    ]
    for i, (label, width) in enumerate(cols2, 1):
        _set_header(ws2, i, label, width)
    ws2.row_dimensions[1].height = 30

    for row_i, lead in enumerate(leads, 2):
        temp = lead.get("temperature", "WARM")
        ts = lead.get("captured_at", "")
        if hasattr(ts, "strftime"):
            ts = ts.strftime("%Y-%m-%d %H:%M")
        note_parts = []
        if temp:               note_parts.append(f"Temperature: {temp}")
        if lead.get("next_step"): note_parts.append(f"Next Step: {lead['next_step']}")
        if lead.get("notes"):  note_parts.append(f"Notes: {lead['notes']}")
        if lead.get("captured_by"): note_parts.append(f"Captured by: {lead['captured_by']}")
        if ts:                 note_parts.append(f"Captured at: {ts}")

        values = [
            lead.get("name", ""),    lead.get("email", ""),
            lead.get("phone", ""),   lead.get("company", ""),
            lead.get("job_title", ""), " | ".join(note_parts),
        ]
        fill = _TEMP_FILL.get(temp)
        for col_i, val in enumerate(values, 1):
            cell = ws2.cell(row=row_i, column=col_i, value=val)
            cell.border = _BORDER
            cell.alignment = _LEFT
            if fill:
                cell.fill = fill
        ws2.row_dimensions[row_i].height = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
